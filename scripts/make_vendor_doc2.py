import json
import os
import re
import psycopg2
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DB = os.environ['DATABASE_URL']
ENRICH_PATH = os.path.join(os.path.dirname(__file__), 'out', 'contact_merged.json')
OUT_PATH = r'C:\Users\sales\Downloads\lr-main\vendors.docx'

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
    rPr.append(color); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

def norm(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())

# --- load data ---
enrich = {}
if os.path.exists(ENRICH_PATH):
    for r in json.load(open(ENRICH_PATH)):
        enrich[r['id']] = r

conn = psycopg2.connect(DB, connect_timeout=15)
cur = conn.cursor()
cur.execute("""
    SELECT id, vendor_type, station, name, address, location, city, phone, website
    FROM vendors
""")
rows = cur.fetchall()

# --- merge: db value wins; fall back to enriched; dedupe identical listings ---
by_key = {}
for vid, vtype, station, name, addr, loc, city, phone, web in rows:
    e = enrich.get(vid, {})
    phone = phone or e.get('phone')
    web = web or e.get('website')
    addr_full = (addr or loc or '').strip()
    key = (norm(name), norm(addr_full))
    if key in by_key:
        old = by_key[key]
        by_key[key] = (old[0], old[1] or station, old[2], old[3], old[4] or phone, old[5] or web)
        continue
    by_key[key] = (vtype or 'other', station, name, addr_full, phone, web)

CAT_LABEL = {
    'restaurant': 'Restaurants', 'bar': 'Bars & Nightlife', 'cafe': 'Cafes & Coffee',
    'beauty': 'Beauty & Barber', 'boutique': 'Boutiques & Shopping',
    'services': 'Services', 'other': 'Other',
}
CAT_ORDER = ['restaurant', 'bar', 'cafe', 'beauty', 'boutique', 'services', 'other']

groups = {}
for vtype, station, name, addr, phone, web in by_key.values():
    groups.setdefault(vtype, {}).setdefault(station or 'No rail stop assigned', []).append(
        (name, addr, phone, web))

for st in groups.values():
    for s in st.values():
        s.sort(key=lambda x: (x[0] or '').lower())

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

doc.add_heading('Light Rail Deals — Vendor Directory', level=0)

first_cat = True
total = 0
missing_phone = missing_web = 0
for vt in CAT_ORDER + sorted(set(groups) - set(CAT_ORDER)):
    if vt not in groups:
        continue
    doc.add_page_break() if not first_cat else None
    first_cat = False
    doc.add_heading(CAT_LABEL.get(vt, vt.title()), level=1)

    def station_sort(s):
        if s == 'No rail stop assigned':
            return (1, s)
        return (0, s.lower())

    for station in sorted(groups[vt], key=station_sort):
        doc.add_heading(station, level=2)
        for name, addr, phone, web in groups[vt][station]:
            total += 1
            p = doc.add_paragraph()
            r = p.add_run(name or 'Unnamed')
            r.bold = True
            p.paragraph_format.space_after = Pt(0)
            if addr:
                pa = doc.add_paragraph(addr)
                pa.paragraph_format.space_after = Pt(0)
            if phone:
                digits = re.sub(r'\D', '', phone)
                pp = doc.add_paragraph()
                pp.paragraph_format.space_after = Pt(0)
                if len(digits) == 10:
                    digits = '1' + digits
                display = phone.strip()
                if len(digits) == 11 and digits.startswith('1'):
                    display = '(%s) %s-%s' % (digits[1:4], digits[4:7], digits[7:])
                add_hyperlink(pp, 'tel:+' + digits, display)
            else:
                missing_phone += 1
                pm = doc.add_paragraph('(phone not found)')
                pm.paragraph_format.space_after = Pt(0)
            if web:
                url = web.strip()
                if not url.lower().startswith(('http://', 'https://')):
                    url = 'https://' + url
                pw = doc.add_paragraph()
                pw.paragraph_format.space_after = Pt(0)
                add_hyperlink(pw, url, web.strip())
            else:
                missing_web += 1
                pw = doc.add_paragraph('(website not found)')
                pw.paragraph_format.space_after = Pt(0)
            doc.add_paragraph()

doc.save(OUT_PATH)
print('wrote', OUT_PATH)
print('unique vendors:', total, '| still missing phone:', missing_phone, '| still missing website:', missing_web)
