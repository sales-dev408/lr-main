import psycopg2
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DB = 'postgresql://postgres:ddfdskhkjdshlkfjhasdlkfjhadskljfhdskljfhdsklfjhadskljfbvdfvuy@db.okbolfpndeakmchpznmt.supabase.co:5432/postgres'

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
    rPr.append(color); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

conn = psycopg2.connect(DB, connect_timeout=15)
cur = conn.cursor()
cur.execute("""
    SELECT vendor_type, station, name, address, phone, website
    FROM vendors
    ORDER BY vendor_type, station, name
""")
rows = cur.fetchall()

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

CAT_LABEL = {
    'restaurant': 'Restaurants', 'bar': 'Bars & Nightlife', 'cafe': 'Cafes & Coffee',
    'beauty': 'Beauty & Barber', 'boutique': 'Boutiques & Shopping',
    'services': 'Services', 'other': 'Other',
}
CAT_ORDER = ['restaurant', 'bar', 'cafe', 'beauty', 'boutique', 'services', 'other']

groups = {}
for vt, station, name, addr, phone, web in rows:
    groups.setdefault(vt or 'other', {}).setdefault(station or 'Unknown Stop', []).append(
        (name, addr, phone, web))

first_cat = True
for vt in CAT_ORDER + sorted(set(groups) - set(CAT_ORDER)):
    if vt not in groups:
        continue
    if not first_cat:
        doc.add_page_break()
    first_cat = False
    doc.add_heading(CAT_LABEL.get(vt, vt.title()), level=1)

    for station in sorted(groups[vt]):
        doc.add_heading(station, level=2)
        for name, addr, phone, web in groups[vt][station]:
            p = doc.add_paragraph()
            r = p.add_run(name or 'Unnamed')
            r.bold = True
            p.paragraph_format.space_after = Pt(0)
            if addr:
                pa = doc.add_paragraph(addr)
                pa.paragraph_format.space_after = Pt(0)
            if phone:
                pp = doc.add_paragraph(phone)
                pp.paragraph_format.space_after = Pt(0)
            if web:
                url = web.strip()
                if not url.lower().startswith(('http://', 'https://')):
                    url = 'https://' + url
                pw = doc.add_paragraph()
                pw.paragraph_format.space_after = Pt(0)
                add_hyperlink(pw, url, web.strip())
            doc.add_paragraph()  # spacer between blocks

out = r'C:\Users\sales\Downloads\lr-main\vendors.docx'
doc.save(out)
print('wrote', out, 'vendors:', len(rows))
